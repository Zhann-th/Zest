"""
pptx_io.py — Модуль ввода/вывода PPTX-файлов.

Отвечает за чтение, модификацию, создание и сохранение PowerPoint-презентаций.
Использует python-pptx для round-trip операций с Open XML.
"""

import os
import re
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))




THEMES = {
    "modern_dark": {
        "bg_color": RGBColor(15, 23, 42),        # #0F172A slate 900
        "card_bg": RGBColor(30, 41, 59),          # #1E293B slate 800
        "title_color": RGBColor(56, 189, 248),    # #38BDF8 sky 400
        "accent_color": RGBColor(129, 140, 248),  # #818CF8 indigo 400
        "text_color": RGBColor(241, 245, 249),    # #F1F5F9 slate 100
        "subtext_color": RGBColor(148, 163, 184), # #94A3B8 slate 400
    },
    "corporate_navy": {
        "bg_color": RGBColor(11, 25, 44),
        "card_bg": RGBColor(30, 62, 98),
        "title_color": RGBColor(245, 158, 11),    # amber
        "accent_color": RGBColor(59, 130, 246),   # blue
        "text_color": RGBColor(248, 250, 252),
        "subtext_color": RGBColor(203, 213, 225),
    },
    "vibrant_gradient": {
        "bg_color": RGBColor(24, 24, 27),         # zinc 900
        "card_bg": RGBColor(39, 39, 42),
        "title_color": RGBColor(236, 72, 153),    # pink
        "accent_color": RGBColor(168, 85, 247),   # purple
        "text_color": RGBColor(255, 255, 255),
        "subtext_color": RGBColor(161, 161, 170),
    },
    "clean_light": {
        "bg_color": RGBColor(248, 250, 252),      # slate 50
        "card_bg": RGBColor(255, 255, 255),
        "title_color": RGBColor(79, 70, 229),     # indigo
        "accent_color": RGBColor(14, 165, 233),   # sky
        "text_color": RGBColor(15, 23, 42),
        "subtext_color": RGBColor(71, 85, 105),
    },
}


class PptxIO:
    """Класс для чтения, записи и модификации PPTX-файлов."""




    def read_pptx(self, filepath):
        """
        Читает PPTX-файл и возвращает структурированные данные.

        Returns:
            dict с ключами:
                - slide_count: int
                - slides: list[dict]
        """
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"Файл не найден: {filepath}")

        prs = Presentation(filepath)
        slides_data = []

        for idx, slide in enumerate(prs.slides):
            slide_info = {
                "slide_index": idx + 1,
                "title": "",
                "subtitle": "",
                "shapes": [],
                "notes": "",
            }

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                slide_info["notes"] = slide.notes_slide.notes_text_frame.text.strip()

            for s_idx, shape in enumerate(slide.shapes):
                if not shape.has_text_frame:
                    continue

                text_frame = shape.text_frame
                full_text = text_frame.text.strip()
                if not full_text:
                    continue

                shape_info = {
                    "shape_index": s_idx,
                    "text": full_text
                }
                slide_info["shapes"].append(shape_info)

                is_title_shape = (shape == slide.shapes.title)
                if is_title_shape or (not slide_info["title"] and len(full_text) < 120):
                    if not slide_info["title"]:
                        slide_info["title"] = full_text

            slides_data.append(slide_info)

        return {
            "slide_count": len(prs.slides),
            "slides": slides_data,
            "filepath": filepath,
        }

    def update_shape_text(self, filepath, slide_index, shape_index, new_text, output_filepath=None):
        """
        Заменяет текст в определенной фигуре на определенном слайде, сохраняя форматирование.
        """
        prs = Presentation(filepath)
        s_idx = slide_index - 1
        
        if s_idx < 0 or s_idx >= len(prs.slides):
            raise IndexError(f"Слайд {slide_index} не существует")
            
        slide = prs.slides[s_idx]
        
        if shape_index < 0 or shape_index >= len(slide.shapes):
            raise IndexError(f"Фигура {shape_index} не существует на слайде {slide_index}")
            
        shape = slide.shapes[shape_index]
        if not shape.has_text_frame:
            raise ValueError("Выбранная фигура не содержит текста")

        tf = shape.text_frame
        for p_idx, paragraph in enumerate(tf.paragraphs):
            for r_idx, run in enumerate(paragraph.runs):
                if p_idx == 0 and r_idx == 0:
                    run.text = new_text
                else:
                    run.text = ""

            if not paragraph.runs and p_idx == 0:
                paragraph.text = new_text
            elif not paragraph.runs:
                paragraph.text = ""
                
        out = output_filepath or filepath
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        prs.save(out)
        return out




    def modify_slide(self, filepath, slide_index, changes, output_filepath=None):
        """
        Модифицирует конкретный слайд в существующем PPTX.

        Args:
            filepath: путь к PPTX
            slide_index: номер слайда (1-based)
            changes: dict с ключами:
                - new_title: str (новый заголовок)
                - new_bullets: list[str] (новые буллет-пункты)
                - replace_map: dict[str, str] (замена строк)
            output_filepath: путь для сохранения (None = перезаписать)

        Returns:
            str — путь к сохранённому файлу
        """
        prs = Presentation(filepath)
        s_idx = slide_index - 1

        if s_idx < 0 or s_idx >= len(prs.slides):
            raise IndexError(f"Слайд {slide_index} не существует (всего {len(prs.slides)})")

        slide = prs.slides[s_idx]
        new_title = changes.get("new_title")
        new_bullets = changes.get("new_bullets")
        replace_map = changes.get("replace_map", {})

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            for search_str, rep_str in replace_map.items():
                if search_str in tf.text:
                    for p in tf.paragraphs:
                        if search_str in p.text:
                            for run in p.runs:
                                if search_str in run.text:
                                    run.text = run.text.replace(search_str, rep_str)

        if new_title:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                for run in slide.shapes.title.text_frame.paragraphs[0].runs:
                    run.text = ""
                slide.shapes.title.text_frame.paragraphs[0].runs[0].text = new_title if slide.shapes.title.text_frame.paragraphs[0].runs else None
                if not slide.shapes.title.text_frame.paragraphs[0].runs:
                    slide.shapes.title.text_frame.paragraphs[0].text = new_title
            else:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        shape.text_frame.paragraphs[0].text = new_title
                        break

        if new_bullets is not None:
            body_shape = None
            for shape in slide.shapes:
                if shape.has_text_frame and shape != slide.shapes.title:
                    body_shape = shape
                    break

            if body_shape:
                tf = body_shape.text_frame
                tf.clear()
                for i, b in enumerate(new_bullets):
                    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                    p.text = b
                    p.level = 0

        out = output_filepath or filepath
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        prs.save(out)
        return out

    def add_slide(self, filepath, slide_data, output_filepath=None, theme_name="modern_dark"):
        """
        Добавляет новый слайд в конец существующей презентации.

        Args:
            filepath: путь к PPTX
            slide_data: dict с title, bullets
            output_filepath: путь для сохранения
            theme_name: тема оформления

        Returns:
            str — путь к файлу
        """
        prs = Presentation(filepath)
        theme = THEMES.get(theme_name, THEMES["modern_dark"])
        blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

        new_slide = prs.slides.add_slide(blank_layout)
        self._set_background(new_slide, theme["bg_color"])

        title_box = new_slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", "Новый слайд")
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = theme["title_color"]
        p.font.name = "Arial"

        bullets = slide_data.get("bullets", [])
        if bullets:
            content_box = new_slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(4.5))
            ctf = content_box.text_frame
            ctf.word_wrap = True
            for i, b in enumerate(bullets):
                bp = ctf.add_paragraph() if i > 0 else ctf.paragraphs[0]
                bp.text = f"• {b}"
                bp.font.size = Pt(20)
                bp.font.color.rgb = theme["text_color"]
                bp.space_before = Pt(12)

        out = output_filepath or filepath
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        prs.save(out)
        return out

    def delete_slide(self, filepath, slide_index, output_filepath=None):
        """
        Удаляет слайд по индексу (1-based).

        Returns:
            str — путь к файлу
        """
        prs = Presentation(filepath)
        s_idx = slide_index - 1

        if s_idx < 0 or s_idx >= len(prs.slides):
            raise IndexError(f"Слайд {slide_index} не существует (всего {len(prs.slides)})")

        rId = prs.slides._sldIdLst[s_idx].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[s_idx]

        out = output_filepath or filepath
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        prs.save(out)
        return out




    def create_presentation(self, title, subtitle="", slides_content=None,
                            template_filename="random", output_filepath="output.pptx"):
        """
        Создаёт презентацию, копируя структуру из готового .pptx шаблона 
        и заполняя плейсхолдеры текстом, сгенерированным ИИ.
        """
        import random
        if slides_content is None:
            slides_content = []

        templates_dir = os.path.join(BASE_DIR, "templates")

        os.makedirs(templates_dir, exist_ok=True)
        
        if template_filename == "random" or template_filename is None:

            available_templates = [f for f in os.listdir(templates_dir) if f.endswith(".pptx")]
            if available_templates:
                template_filename = random.choice(available_templates)
            else:
                template_filename = "base_template.pptx"

        template_path = os.path.join(templates_dir, template_filename)

        if not os.path.exists(template_path):
            prs = Presentation()
            prs.save(template_path)

        prs = Presentation(template_path)

        title_slide_layout = prs.slide_layouts[0]
        title_slide = prs.slides.add_slide(title_slide_layout)
        
        if title_slide.shapes.title:
            title_slide.shapes.title.text = title
            
        if subtitle:
            for shape in title_slide.placeholders:
                if shape.placeholder_format.idx == 1:
                    shape.text = subtitle
                    break

        content_slide_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        
        for s_data in slides_content:
            new_slide = prs.slides.add_slide(content_slide_layout)

            if new_slide.shapes.title:
                new_slide.shapes.title.text = s_data.get("title", "")

            body_shape = None
            for shape in new_slide.placeholders:
                if shape.placeholder_format.idx == 1:
                    body_shape = shape
                    break
                    
            if body_shape and s_data.get("bullets"):
                tf = body_shape.text_frame
                tf.text = s_data["bullets"][0] if len(s_data["bullets"]) > 0 else ""
                for bullet in s_data["bullets"][1:]:
                    p = tf.add_paragraph()
                    p.text = bullet

            if s_data.get("stat_num") and body_shape:
                p = body_shape.text_frame.add_paragraph()
                p.text = f"\nСТАТИСТИКА: {s_data['stat_num']} - {s_data.get('stat_label', '')}"

        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        prs.save(output_filepath)
        return output_filepath

    def save_pptx(self, prs, output_filepath):
        """Сохраняет объект Presentation в файл."""
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        prs.save(output_filepath)
        return output_filepath
