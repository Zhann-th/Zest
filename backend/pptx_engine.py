import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Theme Definitions (RGB Color Palette)
THEMES = {
    "modern_dark": {
        "bg_color": RGBColor(15, 23, 42),       # #0F172A slate 900
        "card_bg": RGBColor(30, 41, 59),        # #1E293B slate 800
        "title_color": RGBColor(56, 189, 248),   # #38BDF8 sky 400
        "accent_color": RGBColor(129, 140, 248), # #818CF8 indigo 400
        "text_color": RGBColor(241, 245, 249),   # #F1F5F9 slate 100
        "subtext_color": RGBColor(148, 163, 184) # #94A3B8 slate 400
    },
    "corporate_navy": {
        "bg_color": RGBColor(11, 25, 44),       # #0B192C dark navy
        "card_bg": RGBColor(30, 62, 98),        # #1E3E62 mid navy
        "title_color": RGBColor(245, 158, 11),   # #F59E0B amber 500
        "accent_color": RGBColor(59, 130, 246),  # #3B82F6 blue 500
        "text_color": RGBColor(248, 250, 252),   # #F8FAFC white
        "subtext_color": RGBColor(203, 213, 225) # #CBD5E1 light gray
    },
    "vibrant_gradient": {
        "bg_color": RGBColor(24, 24, 27),       # #18181B zinc 900
        "card_bg": RGBColor(39, 39, 42),        # #27272A zinc 800
        "title_color": RGBColor(236, 72, 153),   # #EC4899 pink 500
        "accent_color": RGBColor(168, 85, 247),  # #A855F7 purple 500
        "text_color": RGBColor(255, 255, 255),   # white
        "subtext_color": RGBColor(161, 161, 170) # zinc 400
    },
    "clean_light": {
        "bg_color": RGBColor(248, 250, 252),    # #F8FAFC slate 50
        "card_bg": RGBColor(255, 255, 255),     # white
        "title_color": RGBColor(79, 70, 229),    # #4F46E5 indigo 600
        "accent_color": RGBColor(14, 165, 233),  # #0EA5E9 sky 500
        "text_color": RGBColor(15, 23, 42),      # slate 900
        "subtext_color": RGBColor(71, 85, 105)   # slate 600
    }
}

class PPTXEngine:
    def __init__(self):
        pass

    def inspect_pptx(self, filepath):
        """Reads a PPTX file and extracts slides structure, text, shapes, notes."""
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"File not found: {filepath}")

        prs = Presentation(filepath)
        slides_data = []

        for idx, slide in enumerate(prs.slides):
            slide_info = {
                "slide_index": idx + 1,
                "title": "",
                "subtitle": "",
                "bullets": [],
                "all_text": [],
                "notes": ""
            }

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                slide_info["notes"] = slide.notes_slide.notes_text_frame.text.strip()

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                
                text_frame = shape.text_frame
                full_text = text_frame.text.strip()
                if not full_text:
                    continue

                slide_info["all_text"].append(full_text)

                # Try identifying title
                if shape == slide.shapes.title or (not slide_info["title"] and len(full_text) < 100):
                    if not slide_info["title"]:
                        slide_info["title"] = full_text
                        continue

                # Extract bullet points
                for paragraph in text_frame.paragraphs:
                    p_text = paragraph.text.strip()
                    if p_text and p_text != slide_info["title"]:
                        slide_info["bullets"].append(p_text)

            if not slide_info["title"] and slide_info["all_text"]:
                slide_info["title"] = slide_info["all_text"][0]

            slides_data.append(slide_info)

        return {
            "slide_count": len(prs.slides),
            "slides": slides_data
        }

    def _set_background(self, slide, prs, bg_color):
        """Sets full slide solid background color."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

    def create_presentation(self, title, subtitle="", slides_content=[], theme_name="modern_dark", output_filepath="output.pptx"):
        """Creates a brand new custom-styled presentation."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 widescreen format
        prs.slide_height = Inches(7.5)

        theme = THEMES.get(theme_name, THEMES["modern_dark"])

        # 1. Title Slide
        blank_layout = prs.slide_layouts[6]
        title_slide = prs.slides.add_slide(blank_layout)
        self._set_background(title_slide, prs, theme["bg_color"])

        # Decorative Top Accent Bar
        accent_bar = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = theme["title_color"]
        accent_bar.line.fill.background()

        # Title Box
        title_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = theme["title_color"]
        p.font.name = "Arial"
        p.alignment = PP_ALIGN.LEFT

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(22)
            p2.font.color.rgb = theme["subtext_color"]
            p2.font.name = "Arial"
            p2.space_before = Pt(14)

        # Footer badge
        footer_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.333), Inches(0.5))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = "Generated by Trainable PowerPoint AI Studio"
        fp.font.size = Pt(12)
        fp.font.color.rgb = theme["subtext_color"]

        # 2. Add Content Slides
        for s_data in slides_content:
            s_title = s_data.get("title", "Untitled Slide")
            bullets = s_data.get("bullets", [])
            layout_type = s_data.get("layout", "bullets")
            stat_num = s_data.get("stat_num", "")
            stat_label = s_data.get("stat_label", "")

            slide = prs.slides.add_slide(blank_layout)
            self._set_background(slide, prs, theme["bg_color"])

            # Header Accent line
            hdr_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.12), Inches(0.7))
            hdr_bar.fill.solid()
            hdr_bar.fill.fore_color.rgb = theme["title_color"]
            hdr_bar.line.fill.background()

            # Slide Header Text
            hdr_box = slide.shapes.add_textbox(Inches(1.1), Inches(0.7), Inches(11.0), Inches(1.0))
            htf = hdr_box.text_frame
            htf.word_wrap = True
            hp = htf.paragraphs[0]
            hp.text = s_title
            hp.font.size = Pt(32)
            hp.font.bold = True
            hp.font.color.rgb = theme["title_color"]
            hp.font.name = "Arial"

            if layout_type == "stat_callout" and stat_num:
                # Big Stat Card
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
                card.fill.solid()
                card.fill.fore_color.rgb = theme["card_bg"]
                card.line.color.rgb = theme["title_color"]

                ctf = card.text_frame
                ctf.word_wrap = True
                cp1 = ctf.paragraphs[0]
                cp1.text = stat_num
                cp1.font.size = Pt(54)
                cp1.font.bold = True
                cp1.font.color.rgb = theme["title_color"]
                cp1.alignment = PP_ALIGN.CENTER
                
                cp2 = ctf.add_paragraph()
                cp2.text = stat_label
                cp2.font.size = Pt(18)
                cp2.font.color.rgb = theme["text_color"]
                cp2.alignment = PP_ALIGN.CENTER
                cp2.space_before = Pt(12)

                # Right column text
                r_box = slide.shapes.add_textbox(Inches(6.5), Inches(2.0), Inches(6.0), Inches(4.5))
                rtf = r_box.text_frame
                rtf.word_wrap = True
                for b in bullets:
                    rp = rtf.add_paragraph()
                    rp.text = f"• {b}"
                    rp.font.size = Pt(18)
                    rp.font.color.rgb = theme["text_color"]
                    rp.space_before = Pt(10)

            elif layout_type == "two_column":
                half = len(bullets) // 2 or 1
                col1_bullets = bullets[:half]
                col2_bullets = bullets[half:]

                # Col 1
                c1_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.3), Inches(4.5))
                c1_tf = c1_box.text_frame
                c1_tf.word_wrap = True
                for b in col1_bullets:
                    p = c1_tf.add_paragraph()
                    p.text = f"▪ {b}"
                    p.font.size = Pt(18)
                    p.font.color.rgb = theme["text_color"]
                    p.space_before = Pt(12)

                # Col 2
                c2_box = slide.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.3), Inches(4.5))
                c2_tf = c2_box.text_frame
                c2_tf.word_wrap = True
                for b in col2_bullets:
                    p = c2_tf.add_paragraph()
                    p.text = f"▪ {b}"
                    p.font.size = Pt(18)
                    p.font.color.rgb = theme["text_color"]
                    p.space_before = Pt(12)

            else:
                # Standard Bullet Cards layout
                content_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.0), Inches(4.8))
                ctf = content_box.text_frame
                ctf.word_wrap = True

                for i, b in enumerate(bullets):
                    bp = ctf.add_paragraph() if i > 0 else ctf.paragraphs[0]
                    bp.text = f"• {b}"
                    bp.font.size = Pt(20)
                    bp.font.color.rgb = theme["text_color"]
                    bp.font.name = "Arial"
                    bp.space_before = Pt(14)

        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        prs.save(output_filepath)
        return output_filepath

    def modify_existing_pptx(self, filepath, modifications, output_filepath):
        """Applies edits to an existing PowerPoint presentation file."""
        prs = Presentation(filepath)

        for mod in modifications:
            s_idx = mod.get("slide_index", 1) - 1
            if s_idx < 0 or s_idx >= len(prs.slides):
                continue

            slide = prs.slides[s_idx]
            new_title = mod.get("new_title")
            new_bullets = mod.get("new_bullets")
            replace_text_map = mod.get("replace_map", {})

            # Replace explicit text string mappings
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                tf = shape.text_frame

                for search_str, rep_str in replace_text_map.items():
                    if search_str in tf.text:
                        for p in tf.paragraphs:
                            if search_str in p.text:
                                p.text = p.text.replace(search_str, rep_str)

            # Update Title if requested
            if new_title:
                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    slide.shapes.title.text_frame.text = new_title
                else:
                    # Look for first shape
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            shape.text_frame.paragraphs[0].text = new_title
                            break

            # Update bullets if requested
            if new_bullets is not None:
                # Find body shape
                body_shape = None
                for shape in slide.shapes:
                    if shape.has_text_frame and shape != slide.shapes.title:
                        body_shape = shape
                        break
                
                if body_shape:
                    tf = body_shape.text_frame
                    tf.clear()
                    for idx, b in enumerate(new_bullets):
                        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
                        p.text = b
                        p.level = 0

        # Append new slides if requested in modifications
        append_slides = [m for m in modifications if m.get("action") == "add_slide"]
        if append_slides:
            blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
            for s_data in append_slides:
                new_slide = prs.slides.add_slide(blank_layout)
                title_box = new_slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
                title_box.text_frame.paragraphs[0].text = s_data.get("title", "New Slide")
                title_box.text_frame.paragraphs[0].font.size = Pt(32)
                title_box.text_frame.paragraphs[0].font.bold = True

                content_box = new_slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(4.5))
                ctf = content_box.text_frame
                for idx, b in enumerate(s_data.get("bullets", [])):
                    p = ctf.add_paragraph() if idx > 0 else ctf.paragraphs[0]
                    p.text = f"• {b}"
                    p.font.size = Pt(20)

        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        prs.save(output_filepath)
        return output_filepath
